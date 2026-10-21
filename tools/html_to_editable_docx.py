from pathlib import Path
import argparse
import re

from lxml import html
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
FILES = [
    ("Entregable_1_Diagnostico_Informacion_Institucional.html", "Entregable_1_Diagnostico_Informacion_Institucional.docx"),
    ("Entregable_2_Modelo_Informacion_Institucional.html", "Entregable_2_Modelo_Informacion_Institucional.docx"),
    ("Entregable_4_Mapa_Procesos_y_Flujos_Informacion.html", "Entregable_4_Mapa_Procesos_y_Flujos_Informacion.docx"),
]
SOURCE_DIR = ROOT / "entregables" / "fuentes"
WORD_DIR = ROOT / "entregables" / "word"


def clean_text(value):
    return re.sub(r"\s+", " ", value or "").strip()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_bottom_border(cell, color="69B3E7", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def format_footer(paragraph):
    paragraph.text = "Secretaría de Cultura del Estado de Guanajuato  •  PULSO Q  •  Documento editable"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.font.name = "Helvetica Neue"
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(99, 102, 106)


def add_rich_runs(paragraph, node, bold=False, italic=False, mono=False):
    tag = node.tag.lower() if isinstance(node.tag, str) else ""
    bold = bold or tag in {"strong", "b", "th"}
    italic = italic or tag in {"em", "i"}
    mono = mono or "mono" in (node.get("class") or "").split()
    if node.text:
        run = paragraph.add_run(node.text)
        run.bold, run.italic = bold, italic
        if mono:
            run.font.name = "Consolas"
            run.font.size = Pt(8)
    for child in node:
        if isinstance(child.tag, str) and child.tag.lower() == "br":
            paragraph.add_run().add_break()
        else:
            add_rich_runs(paragraph, child, bold, italic, mono)
        if child.tail:
            run = paragraph.add_run(child.tail)
            run.bold, run.italic = bold, italic
            if mono:
                run.font.name = "Consolas"
                run.font.size = Pt(8)


class Converter:
    def __init__(self, source):
        self.source = source
        self.doc = Document()
        section = self.doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.95)
        section.bottom_margin = Inches(0.62)
        section.left_margin = Inches(0.62)
        section.right_margin = Inches(0.62)
        section.header_distance = Inches(0.18)
        section.different_first_page_header_footer = True
        normal = self.doc.styles["Normal"]
        normal.font.name = "Helvetica Neue"
        normal.font.size = Pt(9)
        normal.paragraph_format.space_after = Pt(3)
        normal.paragraph_format.line_spacing = 1.05
        for name, size, color in [("Title", 25, "004B87"), ("Heading 1", 17, "004B87"), ("Heading 2", 14, "004B87"), ("Heading 3", 11, "004B87")]:
            style = self.doc.styles[name]
            style.font.name = "Helvetica Neue"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(color)
            style.font.bold = True

        header = section.header
        table = header.add_table(rows=1, cols=2, width=Inches(7.26))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        left, right = table.rows[0].cells
        left.width = Inches(4.7)
        right.width = Inches(2.56)
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        logo_path = ROOT / "assets" / "identidad" / "Logo.png"
        lp = left.paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        lp.add_run().add_picture(str(logo_path), height=Inches(0.69))
        dependency = lp.add_run("   SECRETARÍA DE CULTURA")
        dependency.font.name = "Helvetica Neue"
        dependency.font.size = Pt(7.5)
        dependency.font.bold = True
        dependency.font.color.rgb = RGBColor.from_string("004B87")
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.paragraph_format.space_after = Pt(0)
        area = rp.add_run("Dirección de Planeación")
        area.font.name = "Helvetica Neue"
        area.font.size = Pt(9)
        area.font.bold = True
        area.font.color.rgb = RGBColor.from_string("004B87")
        set_cell_bottom_border(left)
        set_cell_bottom_border(right)

        format_footer(section.footer.paragraphs[0])
        format_footer(section.first_page_footer.paragraphs[0])

    def add_picture(self, node):
        src = node.get("src")
        if not src:
            return
        path = (self.source.parent / src).resolve()
        if not path.exists():
            return
        with Image.open(path) as image:
            w, h = image.size
        alt = clean_text(node.get("alt"))
        is_logo = any(word in alt.lower() for word in ("logo", "pulso", "gobierno"))
        max_w, max_h = ((2.25, 1.75) if is_logo else (7.0, 7.3))
        width = min(max_w, max_h * w / h)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(str(path), width=Inches(width))
        if alt and "logo" not in alt.lower() and "pulso" not in alt.lower() and "gobierno" not in alt.lower():
            cap = self.doc.add_paragraph(alt)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(8)

    def add_table(self, node):
        is_signature = "signature-table" in set((node.get("class") or "").split())
        rows = node.xpath(".//tr")
        matrix = []
        for row in rows:
            cells = row.xpath("./th|./td")
            if cells:
                matrix.append(cells)
        if not matrix:
            return
        cols = max(sum(int(cell.get("colspan") or 1) for cell in row) for row in matrix)
        table = self.doc.add_table(rows=len(matrix), cols=cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        if is_signature:
            for row_index, word_row in enumerate(table.rows):
                row_props = word_row._tr.get_or_add_trPr()
                row_props.append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    row_props.append(OxmlElement("w:tblHeader"))
        for r, source_cells in enumerate(matrix):
            logical_col = 0
            for source_cell in source_cells:
                span = int(source_cell.get("colspan") or 1)
                target = table.cell(r, logical_col)
                if span > 1:
                    target = target.merge(table.cell(r, logical_col + span - 1))
                logical_col += span
                target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                images = source_cell.xpath(".//img")
                target.text = ""
                if images:
                    for image_node in images:
                        src = image_node.get("src")
                        path = (self.source.parent / src).resolve() if src else None
                        if path and path.exists():
                            with Image.open(path) as image:
                                w, h = image.size
                            max_w, max_h = max(1.1, 6.8 / cols), 3.5
                            width = min(max_w, max_h * w / h)
                            picture_p = target.paragraphs[0]
                            picture_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            picture_p.add_run().add_picture(str(path), width=Inches(width))
                    text_p = target.add_paragraph(clean_text(source_cell.text_content()))
                else:
                    target.text = clean_text(source_cell.text_content())
                for paragraph in target.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    for run in paragraph.runs:
                        run.font.name = "Helvetica Neue"
                        run.font.size = Pt(7.2 if cols >= 5 else 8)
                        if source_cell.tag.lower() == "th" or r == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                if source_cell.tag.lower() == "th" or r == 0:
                    set_cell_shading(target, "004B87")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)

    def add_paragraph_node(self, node, style=None, align=None):
        if not clean_text(node.text_content()):
            return
        p = self.doc.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        add_rich_runs(p, node)

    def add_list(self, node, ordered=False, level=0):
        for li in node.xpath("./li"):
            p = self.doc.add_paragraph(style="List Number" if ordered else "List Bullet")
            p.paragraph_format.left_indent = Inches(0.22 * level)
            if li.text:
                run = p.add_run(li.text)
                run.font.size = Pt(9)
            for child in li:
                tag = child.tag.lower() if isinstance(child.tag, str) else ""
                if tag in {"ul", "ol"}:
                    self.add_list(child, tag == "ol", level + 1)
                elif tag == "br":
                    p.add_run().add_break()
                else:
                    add_rich_runs(p, child)
                if child.tail:
                    p.add_run(child.tail)

    def walk(self, node):
        tag = node.tag.lower() if isinstance(node.tag, str) else ""
        classes = set((node.get("class") or "").split())
        if "toolbar-no-print" in classes or "doc-footer" in classes:
            return
        if "area-block" in classes or "safe-page-break" in classes:
            self.doc.add_page_break()
        if tag == "img":
            self.add_picture(node)
            return
        if tag == "table":
            if "signature-table" in classes and not self.source.name.startswith("Entregable_1_"):
                signature_title = self.doc.add_paragraph("Constancia de entrega para revisión", style="Heading 1")
                signature_title.paragraph_format.page_break_before = True
            self.add_table(node)
            return
        if tag in {"h1", "h2", "h3", "h4"}:
            heading_text = clean_text(node.text_content())
            if heading_text.startswith("Anexo C."):
                self.doc.add_page_break()
            style = {"h1": "Title", "h2": "Heading 1", "h3": "Heading 2", "h4": "Heading 3"}[tag]
            self.add_paragraph_node(node, style=style)
            return
        if tag == "p":
            self.add_paragraph_node(node)
            return
        if tag in {"ul", "ol"}:
            self.add_list(node, tag == "ol")
            return
        if tag == "hr":
            self.doc.add_paragraph("—" * 32).alignment = WD_ALIGN_PARAGRAPH.CENTER
            return
        if tag == "div" and len(node) == 0 and clean_text(node.text_content()):
            text_value = clean_text(node.text_content())
            if "part-title" in classes:
                if text_value.startswith("Constancia de entrega"):
                    self.doc.add_page_break()
                self.add_paragraph_node(node, style="Heading 1")
                return
            style = "Title" if "doc-title" in classes else None
            p = self.doc.add_paragraph(style=style)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(text_value)
            if classes & {"dep-sup", "cover-kicker"}:
                p.runs[0].bold = True
                p.runs[0].font.color.rgb = RGBColor.from_string("004B87")
            if classes & {"dep-sub", "cover-subtitle", "pulso-caption"}:
                p.runs[0].font.color.rgb = RGBColor.from_string("425466")
            return

        for child in node:
            self.walk(child)
        if classes & {"cover", "frontmatter", "factsheet"}:
            self.doc.add_page_break()

    def convert(self, target):
        tree = html.parse(str(self.source))
        body = tree.getroot().find("body")
        self.walk(body)
        while self.doc.paragraphs and not clean_text(self.doc.paragraphs[-1].text):
            p = self.doc.paragraphs[-1]._element
            p.getparent().remove(p)
        self.doc.core_properties.title = clean_text(tree.xpath("string(//title)"))
        self.doc.core_properties.subject = "Entregable editable para revisión y validación"
        self.doc.core_properties.author = "PULSO Q"
        self.doc.save(target)


parser = argparse.ArgumentParser(description="Convierte las fuentes HTML de los entregables a Word editable.")
parser.add_argument("--only", choices=("1", "2", "4"), help="Convierte únicamente el entregable indicado.")
parser.add_argument("--output-suffix", default="", help="Texto que se agrega antes de la extensión .docx.")
args = parser.parse_args()

selected = FILES
if args.only:
    selected = [pair for pair in FILES if pair[0].startswith(f"Entregable_{args.only}_")]

for html_name, docx_name in selected:
    target_name = f"{Path(docx_name).stem}{args.output_suffix}.docx"
    target = WORD_DIR / target_name
    Converter(SOURCE_DIR / html_name).convert(target)
    print(target)
