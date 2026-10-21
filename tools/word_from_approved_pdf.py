"""Reconstruct an editable Word using the approved PDF as layout authority."""
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tmp' / 'pdf-word-deps'))
from pdf2docx import Converter
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml import parse_xml

source = ROOT / 'entregables/pdf/Entregable_1_Diagnostico_Informacion_Institucional.pdf'
target = ROOT / 'entregables/word/Entregable_1_Diagnostico_Informacion_Institucional_BASE_PDF.docx'
converter = Converter(str(source))
try:
    converter.convert(str(target), start=0, end=None)
finally:
    converter.close()

doc = Document(target)
for paragraph in doc._element.body.iter(qn('w:p')):
    texts = list(paragraph.iter(qn('w:t')))
    value = ''.join(t.text or '' for t in texts)
    if 'Luis' in value or 'LUIS' in value:
        for t in texts:
            t.text = (t.text or '').replace('Ángel', 'Angel').replace('ÁNGEL', 'ANGEL').replace('Álvarez', 'Alvarez').replace('ÁL VAREZ','ALVAREZ')
    if value.strip() in ('Secretaría de Cultura del Estado de Guanajuato · Dirección de Planeación', 'Entregable 1 · PULSO Q'):
        # Retain paragraph geometry but move the repeated footer to a real footer.
        for t in texts:
            t.text = ''

for i, section in enumerate(doc.sections):
    section.footer_distance = Pt(0)
    section.different_first_page_header_footer = False
    if i:
        section.footer.is_linked_to_previous = True
footer = doc.sections[0].footer.paragraphs[0]
footer.text = 'Secretaría de Cultura del Estado de Guanajuato · Dirección de Planeación\tEntregable 1 · PULSO Q'
fmt = footer.paragraph_format
fmt.left_indent = Pt(56.7)
fmt.right_indent = Pt(56.7)
fmt.space_before = Pt(0)
fmt.space_after = Pt(0)
fmt.line_spacing = Pt(9)
fmt.tab_stops.add_tab_stop(Pt(498.6), WD_TAB_ALIGNMENT.RIGHT)
for r in footer.runs:
    r.font.name = 'Arial'
    r.font.size = Pt(7.2)
    r.font.color.rgb = RGBColor.from_string('63666A')
def border(paragraph, edge, color, size='6', space='4'):
    props = paragraph._p.get_or_add_pPr()
    borders = props.find(qn('w:pBdr'))
    if borders is None:
        borders = OxmlElement('w:pBdr'); props.append(borders)
    line = OxmlElement('w:' + edge)
    for key, val in {'val':'single','sz':size,'color':color,'space':space}.items():
        line.set(qn('w:' + key),val)
    borders.append(line)
border(footer, 'top', '97999B', '4', '8')
# Anchor the footer to the physical page, so it does not reflow the approved body.
from copy import deepcopy
footer_content = deepcopy(footer._p)
for node in list(footer_content.iter(qn('w:r'))):
    footer_content.remove(node)
footer_run = OxmlElement('w:r')
footer_run_props = OxmlElement('w:rPr')
font = OxmlElement('w:rFonts'); font.set(qn('w:ascii'),'Arial'); font.set(qn('w:hAnsi'),'Arial')
footer_run_props.append(font)
size = OxmlElement('w:sz'); size.set(qn('w:val'),'14'); footer_run_props.append(size)
color = OxmlElement('w:color'); color.set(qn('w:val'),'63666A'); footer_run_props.append(color)
footer_run.append(footer_run_props)
txt = OxmlElement('w:t'); txt.text = 'Secretaría de Cultura del Estado de Guanajuato · Dirección de Planeación'
footer_run.append(txt); footer_content.append(footer_run)
content_props = footer_content.find(qn('w:pPr'))
for key in ('ind','spacing'):
    item = content_props.find(qn('w:' + key))
    if item is not None: content_props.remove(item)
spacing = OxmlElement('w:spacing')
spacing.set(qn('w:before'),'0'); spacing.set(qn('w:after'),'0')
spacing.set(qn('w:line'),'180'); spacing.set(qn('w:lineRule'),'exact')
content_props.append(spacing)
for b in content_props.findall(qn('w:pBdr')): content_props.remove(b)
footer.clear()
footer.paragraph_format.left_indent = Pt(0)
footer.paragraph_format.right_indent = Pt(0)
footer.paragraph_format.line_spacing = Pt(1)
footer.paragraph_format.space_before = Pt(0)
footer.paragraph_format.space_after = Pt(0)
old_border = footer._p.get_or_add_pPr().find(qn('w:pBdr'))
if old_border is not None: footer._p.get_or_add_pPr().remove(old_border)
shape = parse_xml('''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape id="InstitutionalFooter" type="#_x0000_t202" style="position:absolute;margin-left:56.7pt;margin-top:767pt;width:498.6pt;height:14pt;z-index:1;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" stroked="f" filled="f"><v:textbox inset="0,0,0,0"><w:txbxContent/></v:textbox></v:shape></w:pict></w:r>''')
shape.find('.//' + qn('w:txbxContent')).append(footer_content)
footer._p.append(shape)
right_shape = deepcopy(shape)
right_shape.find('.//{urn:schemas-microsoft-com:vml}shape').set('id','InstitutionalFooterRight')
right_shape.find('.//' + qn('w:t')).text = 'Entregable 1 · PULSO Q'
right_props = right_shape.find('.//' + qn('w:pPr'))
jc = OxmlElement('w:jc'); jc.set(qn('w:val'),'right'); right_props.append(jc)
footer._p.append(right_shape)
rule = parse_xml('''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:line id="FooterRule" style="position:absolute;margin-left:0;margin-top:0;mso-position-horizontal-relative:page;mso-position-vertical-relative:page" from="56.7pt,758pt" to="555.3pt,758pt" strokecolor="#97999B" strokeweight="0.5pt"/></w:pict></w:r>''')
footer._p.append(rule)

# The PDF converter fragmented the two-cell signature heading into nested tables.
# Rebuild those two cells as a native, editable Word row at the same position.
heading_table = doc.tables[-2]
labels = ['Elaboró y Entrega (Prestador de Servicios Profesionales)',
          'Recibe para Revisión y Validación (Dependencia Destinataria)']
for cell, label in zip(heading_table.rows[0].cells, labels):
    cell.text = label
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'),'004B87')
    cell._tc.get_or_add_tcPr().append(shd)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Pt(8)
    p.paragraph_format.right_indent = Pt(8)
    p.paragraph_format.line_spacing = Pt(12)
    for r in p.runs:
        r.font.name = 'Arial'; r.font.size = Pt(9); r.bold = True
        r.font.color.rgb = RGBColor(255,255,255)
for cell in doc.tables[-1].rows[0].cells:
    for p in cell.paragraphs:
        if 'LUIS' in p.text or 'LIC.' in p.text:
            border(p,'top','222222','6','5')
            break
for p in doc.paragraphs:
    if p.text.strip() in ('Diagnóstico general','Diagnóstico por área','Anexos técnicos','Constancia de entrega para revisión'):
        p.paragraph_format.right_indent = Pt(56.7)
        border(p,'bottom','004B87','10','0')
doc.core_properties.title = 'Diagnóstico de información institucional'
doc.core_properties.author = 'Luis Angel Durán Alvarez'
doc.save(target)
print(target)
