from pathlib import Path
from collections import Counter
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
import openpyxl, json, re, unicodedata, hashlib, uuid, sys
from openpyxl.utils import get_column_letter as L
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[2]
TMP=Path(__file__).parent
OUT=ROOT/'outputs/01a157b5-4b5c-7dc2-a3d8-393848e44b3f/normalizacion_marcelina_2026-10-23'
OUT.mkdir(parents=True,exist_ok=True)
SOURCE=ROOT/'archivos_base/resumen mio, DGPA.xlsx'
DB=ROOT/'archivos_base/DB_Sistema_Cultura_GTO_V1.xlsx'
hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [SOURCE,DB]}
w=openpyxl.load_workbook(SOURCE,data_only=True,read_only=True)
db=openpyxl.load_workbook(DB,data_only=True,read_only=True)
sn='programa ene-ago 2026'
s=w[sn]
def norm(x):
    return re.sub(r'\s+',' ',''.join(c for c in unicodedata.normalize('NFKD',str(x or '').lower()) if not unicodedata.combining(c))).strip()
def clean(x):
    return re.sub(r'\s+',' ',str(x)).strip() if x is not None else None
def n(x):return x if isinstance(x,(float,int)) and not isinstance(x,bool) else None
def money(x):
    if n(x) is not None:return x
    if isinstance(x,str) and re.fullmatch(r'\$[\d.]+,\d\d',x.strip()):return float(x.strip().replace('$','').replace('.','').replace(',','.'))
    return None
def records(sheet):
    values=list(db[sheet].values); headers=values[0]
    return [dict(zip(headers,r)) for r in values[1:] if r[0] is not None]
act_db=records('Registro_Actividades')
aligned=records('Cat_Alineacion')
catalog={r['ID_Alineacion']:r for r in aligned}
muni={norm(r['Municipio']):r['Municipio'] for r in records('Cat_Municipio')}
recintos={norm(r['Nombre_Recinto']):r for r in records('Cat_Recintos')}
aliases={'san jose itirbide':'san jose iturbide'}
venue_alias={
    'teatro de la cd. de purisima del rincon':'Teatro de la Ciudad de Purísima',
    'teatro de la ciudad de purisima del rincon':'Teatro de la Ciudad de Purísima',
    'teatro de la cd. de purisima':'Teatro de la Ciudad de Purísima',
    'teatro del bicentenario, leon,gto.':'Teatro del Bicentenario',
}
widths={'Registro_Actividades':66,'Registro_Territorio':12,'Registro_Poblacion':8,'Regitro_Evidencia':5}
sheets={name:{'headers':[db[name].cell(1,i).value for i in range(1,width+1)],'rows':[]} for name,width in widths.items()}
used_ids={str(r['ID_Actividad']) for r in act_db}
used_folios={str(r['Folio']) for r in act_db if r.get('Folio')}
all_keys=set(used_ids)
for name in ['Registro_Territorio','Registro_Poblacion','Regitro_Evidencia']:
    all_keys.update(str(r[sheets[name]['headers'][0]]) for r in records(name))
def genkey(kind,row,suffix=''):
    attempt=0
    while True:
        u=uuid.uuid5(uuid.NAMESPACE_URL,f'{hashes[SOURCE.name]}|{sn}|{row}|{kind}|{suffix}|{attempt}').hex
        key=('ACT-SC'+chr(65+int(u[:2],16)%26)+chr(65+int(u[2:4],16)%26)+f'{int(u[4:12],16)%100000:05d}') if kind=='act' else u[:8]
        if key not in all_keys:
            all_keys.add(key);return key
        attempt+=1
def add(name,obj):sheets[name]['rows'].append([obj.get(h) for h in sheets[name]['headers']])
def code(x):
    m=re.search(r'\d+(?:\.\d+)*',str(x or ''))
    return m.group(0) if m else None
def lookup_code(field,x):
    c=code(x)
    opts={clean(r.get(field)) for r in aligned if r.get(field) and code(r[field])==c}
    return next(iter(opts)) if c and len(opts)==1 else None
months={norm(m):i for i,m in enumerate(['enero','febrero','marzo','abril','mayo','junio','julio','agosto','septiembre','octubre','noviembre','diciembre'],1)}
def parsedate(raw):
    if isinstance(raw,datetime):return raw.date().isoformat(),False
    t=norm(raw); mt=re.search(r'(\d{1,2})\s+(?:de\s+)?([a-z]+)(?:\s+(?:de\s+)?(20\d\d))?',t)
    if not mt or mt[2] not in months:raise ValueError('Fecha no interpretable: '+str(raw))
    inferred=not bool(mt[3]); y=int(mt[3] or '2026')
    return datetime(y,months[mt[2]],int(mt[1])).date().isoformat(),inferred
def date_db(x):return str(x)[:10]
source_rows=[]; excluded=[]; lineage=[]; issues=[]
start=datetime.now(ZoneInfo('America/Mexico_City')).replace(microsecond=0)
assert start.date().isoformat()=='2026-10-23', 'La fecha local del lote debe coincidir con el corte autorizado.'
capture=start
for rownum,rr in enumerate(s.iter_rows(min_row=3,values_only=True),3):
    v={L(i+1):x for i,x in enumerate(rr) if x is not None}
    if not v:continue
    source_rows.append(rownum)
    dt,inferred=parsedate(v.get('H'))
    # Two non-identical candidates already in production. Hold aside; never overwrite.
    candidates=[]
    for old in act_db:
        if date_db(old.get('Fecha_Inicio'))!=dt or 'PRODUCCIÓN' not in str(old.get('Direccion')):continue
        if rownum==4 and 'banda' in norm(old.get('Nombre_Actividad')):candidates.append(old['ID_Actividad'])
        if rownum==3 and 'purisima' in norm(old.get('Nombre_Espacio')):candidates.append(old['ID_Actividad'])
    if candidates:
        excluded.append({'fila':rownum,'evento':clean(v.get('J')),'ids_existentes':candidates,'motivo':'Posible actividad ya capturada. No agregar otra vez sin confirmar.'})
        continue
    aid=genkey('act',rownum)
    while True:
        folio='SC-'+capture.strftime('%y%m%d-%H%M%S')
        stamp=capture.isoformat(timespec='seconds')
        capture+=timedelta(seconds=1)
        if folio not in used_folios:used_folios.add(folio);break
    notes=[]
    def issue(text):notes.append(text)
    if inferred:issue('Año 2026 tomado del título de la hoja; fecha original sin año: '+str(v['H']))
    if dt[:4]!='2026':issue('Se conserva el año explícito de la fuente: '+dt+'; confirmar antes de usar en informes de 2026.')
    event=clean(v.get('J'))
    if not event:
        event='Presentación de '+(clean(v.get('K')) or 'artista por identificar')
        issue('La fuente no tiene título. Se preparó una etiqueta descriptiva a partir del artista.')
    city=muni.get(aliases.get(norm(v.get('Q')),norm(v.get('Q'))))
    if not city and 'purisima' in norm(v.get('P')):
        city=muni['purisima del rincon'];issue('Municipio completado por el nombre explícito de la sede en Purísima.')
    if not city:issue('Municipio pendiente de identificar.')
    raw_program=clean(v.get('T'))
    program=raw_program if raw_program and raw_program not in ['N/A','n/a'] else None
    ptext=norm(v.get('S'))
    selected=None
    if program=='PB3574':selected=83
    elif program=='PB3569':selected=85
    elif program=='QC4140' and ('festival' in ptext or not ptext):selected=5
    elif program=='QC4142' and ('ventan' in ptext or not ptext):selected=10
    if selected is None:issue('Clave de proceso / meta pendiente de confirmar; no se asignó una clave por defecto.')
    if selected in [5,10]:issue('Meta propuesta por programa de festivales/artes vivas; validar selección. Dirección conservada como Producción Artística, no SIN ASIGNAR del catálogo.')
    canonical=catalog.get(selected,{})
    a={'ID_Actividad':aid,'Folio':folio,'Fecha_Captura':stamp,'Clave_Proceso':selected,
       'Clave_Proceso_Txt':program,'Meta':canonical.get('Meta'),'Nombre_Meta':canonical.get('Nombre_Meta'),
       'Area':clean(catalog[83]['Area']),'Direccion':clean(catalog[83]['Direccion']),
       'Unidad_Medida':canonical.get('Unidad_Medida'),'Meta_Anual':canonical.get('Meta_Anual'),
       'Nombre_Actividad':event,'Temporalidad':'Única','Fecha_Inicio':dt,'Fecha_Fin':dt,'Total_Sesiones':1,
       'Formato':'Presencial','Usuario_Registra':'machagoya@guanajuato.gob.mx'}
    # Preserve meaning; an attendance count is not a contract or artists-supported count.
    disc=norm(v.get('O'))
    if disc=='teatro':a.update(Disciplina='Artes escénicas',Subdisciplina='Teatro')
    elif disc=='opera':a.update(Disciplina='Artes escénicas',Subdisciplina='Ópera')
    elif disc=='danza':a.update(Disciplina='Artes escénicas');issue('Subdisciplina de danza pendiente; no se eligió un estilo por el título.')
    elif disc=='musica':a.update(Disciplina='Música');issue('Subdisciplina musical pendiente.')
    elif disc=='multidisciplina':a.update(Disciplina='Multidisciplinario');issue('Subdisciplina multidisciplinaria pendiente.')
    else:issue('Disciplina pendiente o valor de fuente no equivalente: '+str(v.get('O')))
    if clean(v.get('C')):
        a['Tipo_Actividad']='Presentación artística';a['Actividad']='Presentación artística';a['Enfoque_Actividad']='Artística'
    if disc=='conferencia':a['Tipo_Actividad']='Conferencia';a['Actividad']='Conferencia'
    venue=clean(v.get('P'))
    rec=recintos.get(norm(venue_alias.get(norm(venue),venue)))
    if rec:
        a['Nombre_Espacio']=rec['Nombre_Recinto'];a['Tipo_Espacio']=rec['Tipo_Espacio']
    elif venue:a['Especifique_Espacio']=venue;issue('Tipo de espacio pendiente para sede no equivalente al catálogo.')
    else:issue('Sede pendiente.')
    ingress=money(v.get('Y'))
    if ingress is not None:a['Genera_Ingresos']=ingress>0;a['Monto_Ingreso']=ingress
    tariff=clean(v.get('R'))
    if tariff:
        a['Detalle_Tarifas']=tariff
        if norm(tariff)=='sin costo':a['Costo']='Sin costo'
        elif '$' in tariff:
            amounts=[float(x.replace(',','')) for x in re.findall(r'\$(\d+(?:\.\d+)?)',tariff)]
            if amounts:a['Costo']='Con costo';a['Costo_Minimo']=min(amounts);a['Costo_Maximo']=max(amounts)
    obj=lookup_code('Objetivo_PG',v.get('A'))
    if obj:
        a['Objetivo_PG']=obj
        a['Eje_PG']=lookup_code('Eje_PG',code(obj).split('.')[0])
    a['Estrategia_PG']=lookup_code('Estrategia_PG',v.get('B'))
    a['Acciones_PG']=lookup_code('Acciones_PG',v.get('E'))
    if norm(v.get('F')) in ['cooperacion','inclusion']:a['Eje_Tranversal_PEC']='Cooperación' if norm(v.get('F'))=='cooperacion' else 'Inclusión'
    if v.get('D') is not None or v.get('G') is not None:issue('Promesa_NC y/o Acciones_PEC pendientes de la guía de códigos internos. Números originales preservados en el control.')
    # A description retains original aggregates that cannot safely become disjoint age rows.
    desc=[]
    for c,label in [('K','Artista o agrupación'),('I','Horario'),('M','Artistas participantes reportados'),('W','Público meta'),('L','Eje reportado'),('S','Programa reportado'),('N','Origen o alcance reportado'),('V','Responsable reportado'),('AO','Observaciones')]:
        if v.get(c) is not None:desc.append(label+': '+clean(v[c]))
    if v.get('X') is not None:desc.append('Honorarios reportados: '+str(v['X']))
    ages=[(c,label) for c,label in [('AE','adultas mayores'),('AF','adultos mayores'),('AG','jóvenes mujeres'),('AH','jóvenes hombres'),('AI','niñas y adolescentes'),('AJ','niños y adolescentes')] if n(v.get(c)) is not None]
    if ages:desc.append('Desglose etario original (sin redistribuir): '+', '.join(label+' '+str(v[c]) for c,label in ages))
    a['Descripcion_Breve']='. '.join(desc)
    add('Registro_Actividades',a)
    total=n(v.get('AK'));female=n(v.get('AA'));male=n(v.get('AB'))
    t={'ID_Territorio':genkey('territorio',rownum),'ID_Actividad':aid,'Municipio':city,'Cantidad':total,'Cantidad_Acceso':total,'Cantidad_Participacion':0 if total is not None else None}
    add('Registro_Territorio',t)
    if total is None:issue('Total territorial vacío en la fuente; se conserva vacío, no cero.')
    if female is not None and male is not None and total is not None and female+male!=total:issue(f'Diferencia conservada: mujeres {female} + hombres {male} = {female+male}; total territorial {total}.')
    if total==0:issue('Total reportado en cero; confirmar si es dato definitivo o pendiente.')
    # Split only gender and disability, never add a subgroup on top of the same total.
    for sex,sexcol,discol in [('Mujer','AA','AC'),('Hombre','AB','AD')]:
        sexcount=n(v.get(sexcol));disabled=n(v.get(discol))
        if sexcount is None:continue
        parts=[]
        if disabled is not None and 0<disabled<=sexcount:
            parts.append(('Personas con Discapacidad',disabled))
            if sexcount-disabled>0:parts.append((None,sexcount-disabled))
        else:
            parts.append((None,sexcount))
            if disabled is not None and disabled>sexcount:issue(f'Discapacidad {sex}: {disabled} supera el conteo por género {sexcount}; cifra conservada en control, sin duplicar población.')
        for pi,(priority,qty) in enumerate(parts):
            add('Registro_Poblacion',{'ID_Poblacion':genkey('poblacion',rownum,sex+str(pi)),
                'ID_Actividad':aid,'Grupo_Etario':'Población General','Genero':sex,'Grupo_Prioritario':priority,
                'Cantidad_Participacion':0,'Cantidad_Acceso':qty,'Cantidad_Poblacion':qty})
    if female is None and male is None and total is not None:
        add('Registro_Poblacion',{'ID_Poblacion':genkey('poblacion',rownum,'global'),'ID_Actividad':aid,
            'Grupo_Etario':'Población General','Genero':'Dato global','Cantidad_Participacion':0,'Cantidad_Acceso':total,'Cantidad_Poblacion':total})
    if female is None and male is None and total is None:issue('Sin cifras de población; no se creó una fila ficticia.')
    if any(n(v.get(c)) is not None for c in ['AE','AF','AG','AH','AI','AJ']):issue('Población cargada como Población General: rangos de edad de origen no confirmados. Desglose original conservado en descripción y control.')
    issue('Público asistente colocado en Acceso; artistas no se sumaron como público ni como agentes contratados. Validar criterio operativo.')
    lineage.append({'fila_fuente':rownum,'id_actividad':aid,'folio':folio,'fecha':dt,'evento':event,'fila_salida':len(sheets['Registro_Actividades']['rows'])+1,'datos_fuente':v,'observaciones':notes})
    for note in notes:issues.append({'fila_fuente':rownum,'ID_Actividad':aid,'observacion':note})

ids=[r[0] for r in sheets['Registro_Actividades']['rows']]
assert len(ids)==len(set(ids)) and not (set(ids)&used_ids)
for name in ['Registro_Territorio','Registro_Poblacion','Regitro_Evidencia']:
    rr=sheets[name]['rows'];assert len(rr)==len({r[0] for r in rr})
    assert all(r[1] in ids for r in rr)
assert len(source_rows)==len(ids)+len(excluded)==244
for p in [SOURCE,DB]:assert hashes[p.name]==hashlib.sha256(p.read_bytes()).hexdigest()
summary={'origen_filas':len(source_rows),'nuevas_actividades':len(ids),'excluidas_por_posible_coincidencia':len(excluded),
 'conteos':{name:len(data['rows']) for name,data in sheets.items()},
 'actividades_sin_clave_proceso':sum(r[sheets['Registro_Actividades']['headers'].index('Clave_Proceso')] is None for r in sheets['Registro_Actividades']['rows']),
 'total_territorial_reportado':sum(r[9] for r in sheets['Registro_Territorio']['rows'] if n(r[9]) is not None),
 'total_poblacion_preparada':sum(r[7] for r in sheets['Registro_Poblacion']['rows']),
 'usuario':'machagoya@guanajuato.gob.mx','preparacion':start.isoformat(),'sha256_fuentes':hashes}
payload={'out':str(OUT),'sheets':sheets,'summary':summary,'lineage':lineage,'excluded':excluded}
(TMP/'carga.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,indent=2))
