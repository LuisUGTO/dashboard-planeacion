from pathlib import Path
import openpyxl, json, datetime, sys
sys.stdout.reconfigure(encoding='utf-8')
from collections import Counter
from openpyxl.utils import get_column_letter as L

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).parent
result={}
for filename in ['DB_Sistema_Cultura_GTO_V1.xlsx','resumen mio, DGPA.xlsx']:
    path=ROOT/'archivos_base'/filename
    w=openpyxl.load_workbook(path,data_only=True)
    f=openpyxl.load_workbook(path,data_only=False)
    book={}
    print('\nFILE',filename)
    for s in w:
        populated=[]
        for row in s:
            vals={c.column_letter:c.value for c in row if c.value is not None}
            if vals: populated.append({'row':row[0].row,'values':vals})
        header=[]
        for row in list(f[s.title].rows)[:min(5,s.max_row)]:
            for c in row:
                if c.value is not None:
                    header.append({'cell':c.coordinate,'value':c.value,'fill':str(c.fill.fgColor),'font':str(c.font.color)})
        book[s.title]={'rows':s.max_row,'cols':s.max_column,'populated_rows':len(populated),'header_styles':header,'data':populated}
        print(s.title,'dimension',s.max_row,s.max_column,'nonempty',len(populated))
        if not filename.startswith('DB_'):
            print('SAMPLE',json.dumps(populated[:4],ensure_ascii=False,default=str)[:4500])
    result[filename]=book
(OUT/'inventario.json').write_text(json.dumps(result,ensure_ascii=False,default=str,indent=2),encoding='utf-8')
