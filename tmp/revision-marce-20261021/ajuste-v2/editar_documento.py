from pathlib import Path
from docx import Document
from docx.shared import Pt,RGBColor
ROOT=Path(__file__).resolve().parents[3]
source=ROOT/'respaldo/revision_ejecutiva_2026-10-21_antes_ajuste_campos/Informe_Ejecutivo_Captura_y_Calidad_2026-10-21.docx'
doc=Document(source)
for p in list(doc.paragraphs):
    if p.text in ['Alerta técnica independiente de la captura'] or p.text.startswith('Indicador_Oficial presenta '):
        p._element.getparent().remove(p._element)
    elif p.text.startswith('Para cada actividad sumamos las filas '):
        for r in p.runs:
            r.text=r.text.replace(' Foto_Evidencia está vacío en las 576 actividades.','')
method=doc.tables[-1]
method.cell(0,0).text='Componente'
method.cell(0,1).text='Contenido del corte'
values=['576 actividades capturadas','565 registros territoriales','1,036 registros de población','291 entradas de evidencia','Catálogo con 99 entradas de metas y alineación']
for i,value in enumerate(values,1):method.cell(i,1).text=value
for i,row in enumerate(method.rows):
    for c in row.cells:
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
            for r in p.runs:r.font.name='Arial';r.font.size=Pt(10);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
target=ROOT/'entregables/revision_ejecutiva_2026-10-21/actualizados/Informe_Ejecutivo_Captura_y_Calidad_2026-10-21.docx'
target.parent.mkdir(parents=True,exist_ok=True)
doc.save(target)
print(target)
