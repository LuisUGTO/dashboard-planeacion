from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'entregables/revision_ejecutiva_2026-10-21'
OUT.mkdir(parents=True,exist_ok=True)
DATA=json.loads((Path(__file__).parent/'resultados.json').read_text(encoding='utf-8'))
assert DATA['conteos']['actividades']==576 and DATA['totales']['difiere']==65
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(1.05);sec.bottom_margin=Inches(.8)
sec.left_margin=sec.right_margin=Inches(.8)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.35)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Arial')
normal=doc.styles['Normal'];normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(24);doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(12)
doc.styles['Subtitle'].font.size=Pt(11)
doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(12);doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].font.bold=True
doc.styles['Heading 2'].paragraph_format.space_before=Pt(8);doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
header=sec.header.paragraphs[0]
header.style=doc.styles['Normal']
header.paragraph_format.tab_stops.clear_all()
header.paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_TAB_ALIGNMENT.RIGHT)
header.add_run().add_picture(str(ROOT/'assets/identidad/Logo.png'),height=Inches(.5))
r=header.add_run('   SECRETARÍA DE CULTURA\tDirección de Planeación');r.font.name='Arial';r.font.size=Pt(9);r.bold=True;r.font.color.rgb=RGBColor.from_string('004B87')
foot=sec.footer.paragraphs[0];foot.style=doc.styles['Normal'];foot.paragraph_format.tab_stops.clear_all();foot.paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_TAB_ALIGNMENT.RIGHT)
r=foot.add_run('Corte al 21 de octubre de 2026   ·   Uso interno de Planeación\t');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('63666A')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
doc.core_properties.author='Luis Angel Durán Alvarez'
doc.core_properties.title='Estado de captura y calidad de la información'
doc.core_properties.subject='Revisión ejecutiva para reuniones de acompañamiento por dirección'

def p(text,bold_lead=None):
    par=doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):par.add_run(bold_lead).bold=True;par.add_run(text[len(bold_lead):])
    else:par.add_run(text)
    return par
def h(text):doc.add_heading(text,1)
def small(text):
    par=p(text)
    for r in par.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('63666A')
    return par
def table(headers,rows,widths,size=10):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'004B87' if ri==0 else ('F1F6FA' if ri%2==0 else 'FFFFFF'));cp.append(sh)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(0);par.paragraph_format.space_before=Pt(0);par.paragraph_format.line_spacing=1.08
                par.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in par.runs:r.font.name='Arial';r.font.size=Pt(size);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

doc.add_paragraph('Estado de captura y calidad\nde la información','Title')
doc.add_paragraph('Acumulado disponible al 21 de octubre de 2026','Subtitle')
p('Para Marcela Chagoya, Dirección de Planeación\nElaboró Luis Angel Durán Alvarez, PULSO Q')
p('Proponemos iniciar el acompañamiento con Museos, Artes Visuales y Exposiciones, por su volumen de pendientes documentales, y Formación e Investigación, por la frecuencia de diferencias entre territorio y población. La convocatoria queda a decisión de Planeación.')
h('Estado general de captura')
p('El corte contiene 576 actividades con identificador único. Seis de las ocho áreas identificadas en el catálogo tienen registros. Forum y Comunicación Social no tienen actividades registradas. Cuatro actividades figuran como Sin asignar y ya están incluidas en las 576.')
table(['Área o categoría','Actividades','Última captura'],[
['Museos, Artes Visuales y Exposiciones',486,'21 oct 2026'],['Formación e Investigación',67,'7 oct 2026'],['Editorial y Publicaciones',9,'6 oct 2026'],['Vinculación y Desarrollo Cultural',5,'21 oct 2026'],['Producción Artística',3,'7 oct 2026'],['Patrimonio Cultural',2,'28 ago 2026'],['Forum',0,'Sin registro'],['Comunicación Social',0,'Sin registro'],['Sin asignar',4,'7 oct 2026']], [4.3,1.0,1.6],10)
small('El 7/9 del dashboard incluye Sin asignar como categoría activa. Para seguimiento institucional equivale a seis de ocho áreas con captura. La ausencia de registros no acredita incumplimiento sin confirmar alcance y obligación de reporte.')

doc.add_page_break();h('Pendientes de calidad por área')
p('Las principales alertas son 310 actividades sin archivo vinculado en el módulo de evidencias, 21 sin población y 15 sin territorio. Entre las 546 actividades con ambas fuentes, 65 presentan totales distintos. Los pendientes se superponen y no deben sumarse como si fueran actividades diferentes.')
table(['Área','Act.','Sin terr.','Sin pob.','Sin evid.','Diferencias'],[
['Museos',486,6,2,273,'8 / 479'],['Formación e Investigación',67,0,6,25,'52 / 61'],['Editorial y Publicaciones',9,0,7,0,'1 / 2'],['Vinculación y Desarrollo',5,5,1,5,'NC'],['Producción Artística',3,2,3,3,'NC'],['Patrimonio Cultural',2,1,1,1,'1 / 1'],['Sin asignar',4,1,1,3,'3 / 3'],['Total',576,15,21,310,'65 / 546']], [2.65,.48,.68,.68,.68,1.73],9.5)
small('Diferencias = actividades con totales distintos / actividades con ambas fuentes. NC = no comparable por ausencia de alguna fuente. Forum y Comunicación Social no aparecen en esta tabla porque no tienen actividades registradas.')
h('Lectura de los resultados')
p('Museos concentra 486 actividades y 273 pendientes de evidencia vinculada. El 56.2% de sus actividades carece de ese vínculo. Formación e Investigación presenta diferencias en 52 de sus 61 actividades comparables, equivalentes al 85.2%, y tiene 25 actividades sin evidencia vinculada.')
p('Editorial y Publicaciones tiene siete de nueve actividades sin población. Las cinco actividades de Vinculación y Desarrollo Cultural carecen de registro territorial. Producción Artística tiene tres actividades sin población ni evidencia vinculada. Cada área debe confirmar si el dato aplica y, en su caso, completarlo o justificarlo.')
h('Integridad observada')
p('No encontramos claves de actividad duplicadas ni registros hijos sin actividad padre en territorio, población o evidencias. Tampoco encontramos cantidades negativas ni diferencias entre participación más acceso y el total de cada fila revisada. Estas verificaciones no sustituyen la validación de los conteos por parte de las áreas.')
small('Un archivo vinculado no equivale a evidencia validada. No se revisaron el contenido, los permisos de apertura ni la suficiencia documental de los archivos.')

doc.add_page_break();h('Reuniones propuestas y casos de trabajo')
doc.add_heading('Museos Artes Visuales y Exposiciones',2)
p('Objetivo: acordar la vinculación de evidencias para 273 actividades y revisar seis actividades sin territorio, dos sin población y ocho diferencias entre fuentes. El número de actividades por sí solo no mide desempeño ni suficiencia de la información.')
doc.add_heading('Formación e Investigación',2)
p('Objetivo: confirmar el criterio de conteo por actividad y de desglose demográfico, revisar las 52 diferencias entre fuentes y establecer cómo atender seis actividades sin población y 25 sin evidencia vinculada.')
doc.add_heading('Aclaraciones adicionales',2)
p('Forum y Comunicación Social requieren confirmar si deben reportar en este sistema, quién es el enlace y cuál es su situación de captura. Patrimonio Cultural requiere aclarar un caso con una diferencia de 49,254 entre fuentes. Planeación debe definir la asignación de cuatro actividades y de tres filas del catálogo con dirección vacía.')
table(['Caso e identificador','Territorio','Población','Diferencia'],[
['Patrimonio Cañada de la Virgen\nACT-SCBV52724','49,465','211','49,254'],['Formación La Cocina del Noreste\nACT-SCWH81648','252','1,512','-1,260'],['Museos Lectura de oráculos griegos\nACT-SCYI34868','96','0','96']], [3.75,1.05,1.05,1.05],9.5)
small('Diferencia = total territorial menos total demográfico. Los ejemplos son alertas, no errores confirmados. Antes de corregir debe comprobarse que ambas fuentes cubran el mismo universo, periodo y criterio de conteo.')
doc.add_heading('Acuerdos sugeridos',2)
p('Cada reunión debe cerrar con un enlace responsable, una lista de actividades por revisar y una fecha de seguimiento. Proponemos revisar un primer corte de atención en cinco días hábiles, sujeto al acuerdo de las áreas. No se propone borrar registros ni modificar producción durante la reunión.')

doc.add_page_break();h('Alcance y metodología del corte')
p('La lectura se realizó el 21 de octubre de 2026 a las 16:33 h de Ciudad de México. Comprende el acumulado disponible, sin filtro por dirección o mes. La base continúa en operación y puede cambiar después de esta lectura.')
table(['Fuente','Registros con identificador'],[
['Registro_Actividades',576],['Registro_Territorio',565],['Registro_Poblacion',1036],['Regitro_Evidencia',291],['Cat_Alineacion',99]],[4.6,2.3],10)
p('Las fechas de captura abarcan del 4 de junio al 21 de octubre de 2026. Las fechas de inicio de las actividades reportadas abarcan del 1 de enero al 11 de octubre de 2026. Algunos registros tienen fecha de término posterior al corte, incluso en 2027. Por ello, las 576 actividades representan registros acumulados y no actividades concluidas.')
p('Para cada actividad sumamos las filas de territorio y población por ID_Actividad, por separado, y comparamos ambos totales. No sumamos las dos fuentes entre sí ni las interpretamos como personas únicas. Una actividad sin evidencia es aquella sin Archivo_Evidencia no vacío en Regitro_Evidencia. Hay 291 registros de evidencia asociados a 266 actividades. Foto_Evidencia está vacío en las 576 actividades.')
doc.add_heading('Alerta técnica independiente de la captura',2)
p('Indicador_Oficial presenta #N/A en las 576 actividades. Nombre_Indicador presenta #ERROR! en 546 y está vacío en 30. Soporte técnico debe revisar estas fórmulas antes de utilizar sus resultados en informes. No atribuimos automáticamente esta situación a los capturistas ni la usamos para seleccionar las dos reuniones propuestas.')
doc.add_heading('Límites de interpretación',2)
p('El análisis no establece obligaciones ni fechas de entrega, no certifica personas únicas y no valida el contenido de las evidencias. Las coincidencias de nombre y fecha no se clasificaron como duplicados confirmados. Las diferencias pueden responder a alcances distintos, desgloses superpuestos o captura pendiente y requieren revisión del área responsable.')
small('Fuente principal: DB_Sistema_Cultura_GTO_V1, tablas indicadas arriba.\nhttps://docs.google.com/spreadsheets/d/1Enq__tSbcolJ7J7PS3qomn5idh9Zh56JdFA9wR9Gkm4/edit\nIdentidad: Manual de Identidad Gobierno de la Gente V1.2, noviembre de 2024, paleta institucional. Consulta de datos en modo de solo lectura.')
for element in [doc._element,doc.styles.element]:
    for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
target=OUT/'Informe_Ejecutivo_Captura_y_Calidad_2026-10-21.docx'
doc.save(target)
print(target)
