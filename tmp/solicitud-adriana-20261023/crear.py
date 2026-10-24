from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'entregables/solicitud_adriana_2026-10-23'
OUT.mkdir(parents=True,exist_ok=True)
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(1.02);sec.bottom_margin=Inches(.68)
sec.left_margin=sec.right_margin=Inches(.8)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Heading 1']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Arial')
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.04
doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 1'].font.size=Pt(12);doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
header=sec.header.paragraphs[0];header.style=doc.styles['Normal']
header.paragraph_format.tab_stops.clear_all()
header.paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_TAB_ALIGNMENT.RIGHT)
pic=header.add_run().add_picture(str(ROOT/'assets/identidad/Logo.png'),height=Inches(.45))
pic._inline.docPr.set('descr','Identidad del Gobierno de Guanajuato')
r=header.add_run('   SECRETARÍA DE CULTURA\tDirección de Planeación');r.font.size=Pt(9);r.bold=True
footer=sec.footer.paragraphs[0];footer.style=doc.styles['Normal']
footer.paragraph_format.tab_stops.clear_all()
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_TAB_ALIGNMENT.RIGHT)
r=footer.add_run('Seguimiento de información de Vinculación\t');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def p(text):return doc.add_paragraph(text)
def h(text):return doc.add_heading(text,1)
def item(lead,text):
    q=doc.add_paragraph();q.add_run(lead+' ').bold=True;q.add_run(text)
    return q

doc.add_paragraph('Propuesta para integrar\nlos reportes de Vinculación','Title')
meta=p('Para Adriana Ugalde · Enero a septiembre de 2026\nLuis Angel Durán Alvarez · 23 de octubre de 2026')
for r in meta.runs:r.font.size=Pt(10)
p('Adriana, gracias por compartir tus reportes. Para atender la carga histórica sin pedirte que captures taller por taller y municipio por municipio, propongo aprovechar la información mensual que ya tienes. El esquema deberá confirmarse con Planeación antes de incorporarlo a la plataforma.')

h('Cómo se propone registrar la información')
p('Preparar un concentrado por mes y por meta o tipo de actividad, conservando sus cantidades reales. Por ejemplo, los 687 talleres de enero seguirán siendo 687 talleres reportados; no se convertirán automáticamente en 687 sesiones ni en una sola actividad realizada.')
p('En los concentrados que agrupen distintas disciplinas, se propone una clasificación general como “Multidisciplinario”. Cuando solo exista el total estatal, se propone “Cobertura estatal”, sin distribuirlo entre municipios ni afirmar que los 46 fueron atendidos. Se conservará cualquier desglose que ya esté disponible, sin pedirte reconstruir información que no tienes.')

h('Qué necesito que me confirmes')
item('1 Periodo y cifras.', 'Confirma que el corte sea enero-septiembre de 2026, qué meses están completos y si los ceros o espacios vacíos significan un dato pendiente. Indica qué versión de las metas anuales está vigente.')
item('2 Talleres y sesiones.', 'Para los demás rubros, aclara qué cuenta cada cantidad: talleres distintos, grupos o sesiones. Si tienes el número de sesiones, compártelo por separado; si no existe, indícalo como no disponible.')
item('3 Población.', 'Confirma si los beneficiarios son personas distintas, inscripciones o asistencias que pueden repetirse. Comparte solo los desgloses por edad, género o municipio que realmente tengas, sin estimarlos.')
item('4 Clasificación y alcance.', 'Confirma las promesas y acciones de los programas de Gobierno y Cultura que ya anotaste, y completa los ejes pendientes. Señala qué rubros agrupan disciplinas y cuáles requieren cobertura estatal. Indica también qué programas y actividades deseas incluir.')
item('5 Agentes culturales.', 'Distingue al público de los instructores, promotores u otros agentes. Indica el tipo y cantidad de agentes, a qué periodo corresponden y si cuentan con contrato con la Secretaría; los asistentes no se registrarán como agentes contratados.')
item('6 Respaldos y registros previos.', 'Comparte los reportes o archivos de respaldo disponibles e indica qué meses o actividades ya capturaste, para revisar coincidencias antes de agregar información.')

h('Cómo continuar')
p('Puedes responder sobre tus archivos actuales, por WhatsApp o mediante un audio. No necesitas llenar las cuatro hojas de la base. Con tus confirmaciones prepararé la propuesta de carga para validarla contigo y con Planeación, y revisaré los ajustes necesarios para que el tablero distinga talleres, sesiones y concentrados mensuales.')

doc.core_properties.author='Luis Angel Durán Alvarez'
doc.core_properties.title='Propuesta para integrar los reportes de Vinculación'
doc.core_properties.subject='Confirmaciones para una carga histórica mensual de enero a septiembre de 2026'
doc.core_properties.language='es-MX'
for el in [doc.styles.element,doc.element,sec.header._element,sec.footer._element]:
    for border in el.xpath('.//w:pBdr'):border.getparent().remove(border)
target=OUT/'Solicitud_Informacion_Adriana_Vinculacion_2026.docx'
doc.save(target)
print(target)
print('Palabras',len(' '.join(q.text for q in doc.paragraphs).split()))
