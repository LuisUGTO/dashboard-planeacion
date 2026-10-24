from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'entregables/solicitud_teatros_2026-10-22'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(1.05), Inches(.8)
sec.left_margin = sec.right_margin = Inches(.8)
sec.header_distance, sec.footer_distance = Inches(.25), Inches(.35)
for name in ['Normal', 'Title', 'Heading 1']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
pf = doc.styles['Normal'].paragraph_format
pf.space_after, pf.line_spacing = Pt(7), 1.08
doc.styles['Title'].font.size = Pt(22)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 1'].font.size = Pt(13)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)

head = sec.header.paragraphs[0]
head.style = doc.styles['Normal']
head.paragraph_format.tab_stops.clear_all()
head.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), WD_TAB_ALIGNMENT.RIGHT)
head.add_run().add_picture(str(ROOT / 'assets/identidad/Logo.png'), height=Inches(.48))
r = head.add_run('   SECRETARÍA DE CULTURA\tIntegración de información')
r.bold = True
r.font.size = Pt(9)
r.font.color.rgb = RGBColor.from_string('004B87')
foot = sec.footer.paragraphs[0]
foot.style = doc.styles['Normal']
foot.paragraph_format.tab_stops.clear_all()
foot.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), WD_TAB_ALIGNMENT.RIGHT)
r = foot.add_run('Solicitud de aclaraciones · Teatros · 22 de octubre de 2026\t')
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string('63666A')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
foot._p.append(field)

def p(text):
    return doc.add_paragraph(text)

def h(text):
    doc.add_heading(text, 1)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for values in rows:
        for c, text in zip(t.add_row().cells, values):
            c.text = text
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + side)
        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            b.set(qn('w:' + key), val)
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for ri, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci, c in enumerate(row.cells):
            c.width = Inches(widths[ci])
            sh = OxmlElement('w:shd')
            sh.set(qn('w:fill'), '004B87' if ri == 0 else ('F1F6FA' if ri % 2 == 0 else 'FFFFFF'))
            c._tc.get_or_add_tcPr().append(sh)
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(5)
                par.paragraph_format.space_before = Pt(5)
                for run in par.runs:
                    run.font.size = Pt(10)
                    run.bold = ri == 0
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if ri == 0 else '000000')
    return t

doc.add_paragraph('Solicitud de aclaraciones\nInformación de Teatros', 'Title')
p('Para: Marcelina · Área de Teatros\nDe: Luis Angel Durán Alvarez · PULSO Q\nFecha: 22 de octubre de 2026')
p('Marcelina, gracias por compartir el archivo “resumen mio, DGPA.xlsx” y la explicación de su numeración. Para revisar si puedo preparar la información para su incorporación a la plataforma, necesito confirmar los siguientes puntos. No es necesario rehacer el archivo; puedes responder por apartado y compartir las guías que ya utilizas.')

h('1. Archivo definitivo, periodo y unidad de registro')
p('¿Qué hoja debo tomar como fuente definitiva: “programa ene-ago 2026”, “julio - sep”, “Hoja3” o “Hoja4”? Hay coincidencias entre hojas; confirma cuáles son cortes, copias o versiones complementarias para no duplicar información.')
p('Indica el periodo y la fecha de corte que se deben incorporar. En “Hoja3” aparecen actividades de noviembre y diciembre: ¿cuáles ya se realizaron, cuáles siguen programadas y cuáles fueron canceladas o reprogramadas? Confirma también si cada fila representa una función, un evento o un conjunto de funciones.')
p('Si ya capturaron parte de esta información en la plataforma, señala qué actividades o periodo. Si tienes sus identificadores, compártelos; si no, se revisarán las coincidencias antes de cualquier carga.')

h('2. Guía de equivalencias de los códigos')
p('Comparte las hojas impresas o fotografías legibles donde relacionas tus números con las opciones de la plataforma. Para cada campo codificado necesitamos el código usado y el texto completo de la opción, distinguiendo la numeración oficial de la que asignaste internamente.')
p('Esto aplica a objetivo, estrategia, tipo de actividad, Promesa del Nuevo Comienzo, Acción del Programa de Gobierno, eje transversal y Acción del Programa Estatal de Cultura, según corresponda. Conserva códigos como “6.1” tal como los usaste y confirma si su significado es el mismo en todas las hojas y meses.')

doc.add_page_break()
h('3. Asistencia y desgloses de población')
p('Confirma cuál es el total válido de asistencia por fila y qué significan los ceros y las celdas vacías: sin asistentes, dato pendiente, no disponible o actividad no realizada. Indica si se cuentan asistencias por función o personas únicas.')
p('En “julio - sep” hay, por ejemplo, estas diferencias. ¿Qué cifra corresponde o por qué representan conteos distintos? Las referencias son las filas originales del Excel.')
table(['Referencia', 'Mujeres + hombres', 'Total en AD'], [
    ['Fila 17 · T17, U17 y AD17', '197 + 164 = 361', '164'],
    ['Fila 19 · T19, U19 y AD19', '1,100 + 890 = 1,990', '0'],
], [3.2, 2.1, 1.6])
p('Precisa los rangos de edad de niñas y adolescentes, jóvenes, personas adultas y adultas mayores, y si se traslapan. Confirma si discapacidad es un subconjunto del público ya contado; no debe sumarse otra vez al total por esa condición.')

h('4. Artistas, territorio y programas')
p('¿“Número de artistas participantes” cuenta personas, agrupaciones o participaciones por función? ¿Una misma persona se repite en distintas funciones? Indica el criterio para distinguir artistas o agentes culturales del público asistente.')
p('¿“NACIONAL / INTERNACIONAL / ESTATAL” describe el origen de los artistas, el alcance del evento u otro concepto? Confirma además la sede y municipio de realización. Esa etiqueta no se interpretará automáticamente como cobertura de los 46 municipios.')
p('En “Presupuesto” aparecen claves como PB3574, QC4140 y QC4142. Confirma qué representan y su nombre completo, para no tratarlas como importes ni asignarlas a un programa incorrecto.')

h('5. Comprobantes, correcciones y validación')
p('Aclara qué representa la columna “Comprobante”, donde aparecen cifras. Si hay evidencias en archivos o Drive, comparte su ubicación y la forma de relacionarlas con cada actividad; no es necesario volver a generar documentos existentes.')
p('Puedes enviar las correcciones indicando hoja, fila, campo y valor correcto. Confirma quién del área validará la información preparada y los totales antes de incorporarla.')

h('Siguiente paso')
p('Con tus respuestas podré revisar la viabilidad, estimar el tiempo necesario y acordar el alcance del apoyo con el área y Planeación. Esta solicitud no implica una carga comprometida: primero se prepararía una muestra para validación y después se acordaría cualquier incorporación a la plataforma.')

doc.core_properties.author = 'Luis Angel Durán Alvarez'
doc.core_properties.title = 'Solicitud de aclaraciones para integrar información de Teatros'
doc.core_properties.subject = 'Revisión previa a una posible normalización y carga asistida'
for element in [doc.styles.element, doc.element, sec.header._element, sec.footer._element]:
    for border in element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
target = OUT / 'Solicitud_Aclaraciones_Teatros_2026-10-22.docx'
doc.save(target)
print(target)
