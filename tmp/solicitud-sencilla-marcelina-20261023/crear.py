from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'entregables/solicitud_marcelina_2026-10-23'
OUT.mkdir(parents=True,exist_ok=True)
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(1.02);sec.bottom_margin=Inches(.68)
sec.left_margin=sec.right_margin=Inches(.8)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Heading 1']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Arial')
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.04
doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 1'].font.size=Pt(12);doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)

header=sec.header.paragraphs[0];header.style=doc.styles['Normal']
header.paragraph_format.tab_stops.clear_all()
header.paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_TAB_ALIGNMENT.RIGHT)
picture=header.add_run().add_picture(str(ROOT/'assets/identidad/Logo.png'),height=Inches(.45))
picture._inline.docPr.set('descr','Identidad del Gobierno de Guanajuato')
r=header.add_run('   SECRETARÍA DE CULTURA\tDirección de Planeación')
r.font.size=Pt(9);r.bold=True
footer=sec.footer.paragraphs[0];footer.style=doc.styles['Normal']
footer.paragraph_format.tab_stops.clear_all()
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_TAB_ALIGNMENT.RIGHT)
r=footer.add_run('Seguimiento de información de Teatros\t');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text):return doc.add_paragraph(text)
def h(text):return doc.add_heading(text,1)
doc.add_paragraph('Información para completar\nel registro de Teatros','Title')
meta=p('Para Marcelina · Enero a agosto de 2026\nLuis Angel Durán Alvarez · 23 de octubre de 2026')
for r in meta.runs:r.font.size=Pt(10)
p('Marcelina, gracias por compartir tu concentrado. Ya avancé en acomodar la información para la plataforma. Antes de incorporarla, necesito tu apoyo para confirmar algunos datos. No hace falta que vuelvas a capturar todo ni que llenes las hojas técnicas de la base.')

h('1 Tus anotaciones y las opciones que elegiste')
p('¿Puedes compartirme fotos legibles de las hojas donde numeraste las opciones, o una lista con el número y su significado? Por ejemplo, qué opción representa el 9 en Promesa del Nuevo Comienzo y el 16.1 en la Acción del Programa Estatal de Cultura.')
p('Para las actividades sin selección, necesito saber qué acción del Programa de Gobierno, acción del Programa Estatal de Cultura, promesa y eje transversal les corresponden. Si varias comparten las mismas opciones, puedes indicarme el grupo y sus excepciones; no es necesario repetirlo en cada renglón. Con esas respuestas podré completar los campos relacionados.')

h('2 Qué representa cada renglón')
p('¿Cada renglón corresponde a una sola función o presentación? Si es así, usaríamos una sesión. Si alguno reúne varias funciones o sesiones, indícame cuáles son y cuántas incluye. Confirma también si todo corresponde a enero-agosto de 2026 o si hay alguna fecha de otro año que deba revisarse.')

h('3 La comunidad o colonia donde se realizó')
p('Cuando no esté claro en el concentrado, necesito el nombre de la comunidad o colonia y su municipio. Si ya aparece en la descripción de la sede, lo revisaré sin pedirte que lo escribas otra vez. Si no cuentas con el dato, puedes indicarlo como pendiente; no es necesario adivinarlo.')

h('4 Los datos que todavía estén pendientes')
p('Por ahora conservaré las cifras que reportaste. Solo te pido identificar si los ceros o espacios vacíos son datos pendientes o cantidades definitivas. Si tienes reportes, comprobantes u otro concentrado con información adicional, puedes compartirlos o indicar dónde están, sin volver a elaborarlos.')

p('Puedes responder por WhatsApp, audio o con anotaciones en tu archivo actual, indicando el evento y la fecha cuando sea un caso específico. Primero me servirían tus hojas de anotaciones. Después te compartiré los casos que sigan pendientes para que valides la información antes de la carga.')
doc.core_properties.author='Luis Angel Durán Alvarez'
doc.core_properties.title='Información para completar el registro de Teatros'
doc.core_properties.subject='Confirmaciones para integrar el concentrado de enero a agosto de 2026'
doc.core_properties.language='es-MX'
for el in [doc.styles.element,doc.element,sec.header._element,sec.footer._element]:
    for border in el.xpath('.//w:pBdr'):border.getparent().remove(border)
target=OUT/'Solicitud_Informacion_Marcelina_Enero_Agosto_2026.docx'
doc.save(target)
print(target)
print('Palabras',len(' '.join(x.text for x in doc.paragraphs).split()))
